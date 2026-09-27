#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成研学论文 Word 文档"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()

# 页面设置
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(3)

def set_font(run, name_cn='宋体', name_en='Times New Roman', size=12, bold=False, color=None):
    run.font.name = name_en
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), name_cn)

def add_title(text, size=18):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_font(run, name_cn='黑体', size=size, bold=True)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_subtitle(text, size=14):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_font(run, name_cn='黑体', size=size, bold=False, color=RGBColor(0x55, 0x55, 0x55))
    p.paragraph_format.space_after = Pt(18)
    return p

def add_h1(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, name_cn='黑体', size=14, bold=True)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_h2(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, name_cn='黑体', size=12, bold=True)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_body(text, author=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    if author:
        run_a = p.add_run(f'【{author}】')
        set_font(run_a, name_cn='楷体', size=10.5, color=RGBColor(0x88, 0x88, 0x88))
    run = p.add_run(text)
    set_font(run, name_cn='宋体', size=12)
    return p

def add_quote(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_font(run, name_cn='楷体', size=12, color=RGBColor(0x33, 0x33, 0x33))
    return p

def add_meta(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_font(run, name_cn='仿宋', size=10.5, color=RGBColor(0x77, 0x77, 0x77))
    p.paragraph_format.space_after = Pt(12)

# ============ 正文开始 ============

add_title('我们在《唐诗三百首》里发现了黄昏')
add_subtitle('——关于唐代诗人"日暮"书写的探究报告')
add_meta('小组研学论文  |  作者：甲、乙')

# 一
add_h1('一、我们是怎么发现这个题目的')
add_body('这次研学活动要求我们从《唐诗三百首》里选一个主题来探究。一开始我们想了很多：写李白和杜甫？写月亮？写送别诗？可这些题目好像别人都做过了。', author='作者甲')
add_body('直到有一天，我们组在自习课上一起背诗，背到崔颢的《黄鹤楼》——"日暮乡关何处是？烟波江上使人愁"，又翻到孟浩然的《宿建德江》——"移舟泊烟渚，日暮客愁新"，接着是李商隐的《乐游原》——"夕阳无限好，只是近黄昏"。', author='作者甲')
add_body('我们忽然发现一件事：怎么这么多诗都在写傍晚？于是我们做了一件笨功夫的事：把《唐诗三百首》从头到尾翻了一遍，把里面写到"日暮""斜阳""落日""向晚""返景"的诗都挑出来。结果挑出了不下十五首，从五言古诗到七言绝句都有。', author='作者甲')
add_body('我们的探究问题就这样产生了：唐代诗人为什么这么喜欢写黄昏？同样是看夕阳，为什么有人写得安静，有人写得难过，有人写得沉重？', author='作者甲')

# 二
add_h1('二、第一个发现：黄昏为什么总是和"愁"绑在一起？')
add_body('带着问题，我们先把挑出来的诗读了一遍。很快发现一个规律：黄昏在唐诗里几乎不是一个"好看"的时间，而是一个"发愁"的时间。', author='作者甲')
add_body('崔颢站在黄鹤楼上，白天看到的是"晴川历历汉阳树，芳草萋萋鹦鹉洲"，明明是很明朗的景色。可一到"日暮"，他突然就"愁"了。孟浩然把船停在小洲边，本来只是歇歇脚，可"日暮"一来，"客愁新"——新的愁绪就冒出来了。', author='作者甲')
add_body('为什么？我们讨论了很久，觉得大概有三个原因：', author='作者甲')
add_body('第一，对比。黄昏的时候，鸡回窝了，牛回圈了，鸟归林了。王维自己写过"斜光照墟落，穷巷牛羊归"——连牛羊都知道回家。可船上的诗人呢？他是"客"，他没有家可回。所有人都在回家的路上，只有他还在漂着。', author='作者甲')
add_body('第二，时间过去了。白天是用来做事的，可一到傍晚，一天就结束了。对离家在外的人来说，"又过了一天"不是一句轻松的话——离家又远了一天，回家又晚了一天。', author='作者甲')
add_body('第三，看不清楚了。日暮时分，江上起雾，远山模糊。"烟波江上"四个字，把"看不清"写成了画面。看不清前路，也看不清故乡，这种迷茫本身就是愁。', author='作者甲')

# 三
add_h1('三、第二个发现：三种黄昏，三种人')
add_body('可问题来了：如果黄昏只是用来写愁的，那为什么我们读王维的黄昏，不觉得愁？于是我们把十五首诗按诗人分组，一组一组地读。读着读着，我们发现了一个更有意思的现象：同样一个夕阳，在不同诗人笔下，味道完全不一样。', author='作者甲')
add_body('我们选出了三位最有代表性的诗人，做了对比。', author='作者甲')

add_h2('3.1 王维的黄昏：光来了，就接住')
add_body('王维写过一首很短的诗，叫《鹿柴》：', author='作者甲')
add_quote('空山不见人，但闻人语响。')
add_quote('返景入深林，复照青苔上。')
add_body('深山里没有人，只有一缕夕阳的余晖斜斜照进来，落在青苔上。就这么一个小小的画面，王维写了二十个字。他没有说"夕阳好美"，也没有说"可惜要天黑了"。他只是把这个瞬间记下来了。', author='作者甲')
add_body('再看《渭川田家》："斜光照墟落，穷巷牛羊归。野老念牧童，倚杖候荆扉。"夕阳金色的光洒在村子上，老爷爷拄着拐杖等孙子回家。这是一个多么普通的傍晚，可王维写得像一幅画。', author='作者甲')
add_body('我们组的结论是：王维面对黄昏，是不发愁的。光来了就看光，天黑了就关门。他在《送别》里写"山中相送罢，日暮掩柴扉"——送走朋友，天暗了，关上门。没有哭，没有感慨，就是安安静静地过日子。', author='作者甲')
add_body('这是第一种态度：安住当下。', author='作者甲')

add_h2('3.2 李商隐的黄昏：在"近"字里停了一下')
add_body('可李商隐不一样。他的《乐游原》大概是整个《唐诗三百首》里最著名的二十个字：', author='作者甲')
add_quote('向晚意不适，驱车登古原。')
add_quote('夕阳无限好，只是近黄昏。')
add_body('我们一开始读，觉得这首诗好简单啊，连中学生都能看懂。可读着读着，我们注意到一个字——"近"。', author='作者甲')
add_body('"夕阳无限好"是赞美。"只是近黄昏"是转折。可李商隐说的不是"夕阳就是黄昏"，而是"夕阳近黄昏"。"近"是什么意思？是还没到，但已经感觉到它要来了。', author='作者甲')
add_body('我们组有人说了一句特别好的话："这不就是周末晚上吗？你正在玩，忽然意识到明天要上学了——那种开心里面掺着一点慌。"我们都笑了，但想想确实是这么回事。李商隐写下的，就是这种"马上就要结束了"的感觉。最美好的时刻，恰恰是因为你知道它快没了，才格外珍贵。', author='作者甲')
add_body('这是第二种态度：清醒地珍惜。', author='作者甲')

add_h2('3.3 杜甫的黄昏：一个人扛起了整个时代')
add_body('读到杜甫的时候，我们组安静了好一会儿。《登高》：', author='作者乙')
add_quote('风急天高猿啸哀，渚清沙白鸟飞回。')
add_quote('无边落木萧萧下，不尽长江滚滚来。')
add_quote('万里悲秋常作客，百年多病独登台。')
add_quote('艰难苦恨繁霜鬓，潦倒新停浊酒杯。')
add_body('这首诗被前人称为"古今七言律第一"。我们读的时候，最直观的感受就是两个字：重。风是急的，天是高的，猿声是哀的，树叶萧萧往下落，江水滚滚往东流。天地那么大，时间那么长，可站在高台上的，是一个五十六岁、生着病、离家万里、独自一个人的老人。', author='作者乙')
add_body('王维的黄昏里有他自己，李商隐的黄昏里也有他自己。可杜甫的黄昏里，除了他自己，还有别人。他在《登楼》里写"花近高楼伤客心，万方多难此登临"——他看到一朵花，想到的不是花好不好看，而是天下正在打仗，老百姓正在受苦。', author='作者乙')
add_body('我们讨论：杜甫是不是太沉重了？后来我们觉得，不是沉重，是担当。他没有逃避自己的愁，也没有只盯着自己的愁。他把个人的遭遇和整个时代连在了一起。', author='作者乙')
add_body('这是第三种态度：把黄昏扛在肩上。', author='作者乙')

# 四
add_h1('四、这个发现，跟我们有什么关系？')
add_body('读到这里，我们一直在讨论：这些唐朝人看夕阳的心情，跟我们有什么关系？', author='作者乙')
add_body('我们是中学生。每天早上六点多起床，背着书包上学，白天上课，晚上写作业写到很晚，周末还有补习班。我们的时间表排得满满的，好像从来没有时间"停下来看夕阳"。', author='作者乙')
add_body('有同学说，他放学路上其实能看到晚霞，但他永远在低头看手机，或者在想明天要交的作业。有同学说，他上一次认真看天，还是小学的时候。', author='作者乙')
add_body('我们忽然明白了一件事：王维、李商隐、杜甫在一千多年前做的事情，其实就是我们现在最缺的——认真对待"停下来"的那一刻。', author='作者乙')
add_body('王维教我们的是：不用想太多，光来了就看看。就像写作业写到很累的时候，抬头看看窗外，不用自责，也不用焦虑，那束光本来就在那里。', author='作者乙')
add_body('李商隐教我们的是：美好的东西会结束，但正因为会结束，才值得好好经历。就像周末、就像假期、就像和朋友待在一起的时间——不用因为"快没了"而扫兴，反而要更认真地过。', author='作者乙')
add_body('杜甫教我们的是：除了自己的烦心事，还可以看看更大的世界。我们有考试的压力，有排名的焦虑，但这些不是全部。知道别人也在努力，知道这个世界很大，个人的那点愁就不会把自己压垮。', author='作者乙')

# 五
add_h1('五、我们的结论')
add_body('回到最开始的问题：唐代诗人为什么那么喜欢写黄昏？', author='作者乙')
add_body('因为黄昏是一天里唯一可以停下来的时刻。白天要做事，晚上要休息，只有傍晚——光还在，但事已经做完了——人会不由自主地抬头，看看天，想想自己。', author='作者乙')
add_body('一千多年过去了，我们的生活变了：我们有手机，有网络，有做不完的卷子。可我们和唐朝诗人面对的是同一个夕阳。', author='作者乙')
add_body('我们这次探究最大的收获，不是背会了几首诗，而是发现了一件事：原来在十几岁的年纪，我们也可以学着像王维那样接一束光，像李商隐那样珍惜一个傍晚，像杜甫那样把自己的生活和更大的世界连起来。', author='作者乙')
add_body('夕阳无限好。不是因为它不会消失。而是因为总有人愿意在它消失之前，停下来，认真看一眼。', author='作者乙')

# 参考文献
add_h1('参考文献')
refs = [
    '[1] 蘅塘退士编. 唐诗三百首[M]. 中华书局.',
    '[2] 马茂元等. 唐诗鉴赏辞典[M]. 上海辞书出版社.',
    '[3] 傅道彬. 晚唐钟声：中国文学的原型批评（修订本）[M]. 北京大学出版社.',
    '[4] 部编版语文教材七年级上册、八年级上册相关篇目.',
]
for r in refs:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(r)
    set_font(run, name_cn='宋体', size=10.5)

out = '/home/user/Doubao/chats/38444192369057538/唐诗研学/我们在唐诗三百首里发现了黄昏_研学论文.docx'
doc.save(out)
print(f'Saved: {out}')
